import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from django.core.exceptions import ValidationError
from django.core.validators import URLValidator, validate_email
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.contractors.models import Company, Contact, TradeCapability
from apps.contractors.providers import ContractorResult
from apps.contractors.services import (
    create_contact,
    dedupe_company,
    normalize_domain,
    normalize_phone,
    normalized_name,
)
from apps.projects.audit import record_event
from apps.scope_packages.trades import TRADE_CHOICES

from .models import ProspectImport, ProspectListEntry, ProspectTag

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
MAX_XLSX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024

CANONICAL_FIELDS = (
    "company_name",
    "website",
    "company_phone",
    "company_email",
    "address",
    "city",
    "province",
    "postal_code",
    "country",
    "trade",
    "contact_name",
    "contact_title",
    "contact_email",
    "contact_phone",
    "tags",
    "notes",
)

IMPORT_TEMPLATE_FILENAME = "BB-Builders-Prospect-Import-Template.xlsx"
IMPORT_TEMPLATE_SAMPLE = (
    "Example Mechanical Ltd",
    "https://example.com",
    "+1 416 555 0100",
    "info@example.com",
    "100 Example Street",
    "Toronto",
    "Ontario",
    "M5V 1A1",
    "Canada",
    "HVAC / Mechanical",
    "John Example",
    "Estimator",
    "john@example.com",
    "+1 416 555 0101",
    "hvac,toronto",
    "Example row - delete before importing",
)


def build_import_template():
    """Return an in-memory XLSX aligned with the canonical import schema."""
    workbook = Workbook()
    prospects = workbook.active
    prospects.title = "Prospects"
    prospects.append(CANONICAL_FIELDS)
    prospects.append(IMPORT_TEMPLATE_SAMPLE)
    prospects.freeze_panes = "A2"
    prospects.auto_filter.ref = f"A1:{get_column_letter(len(CANONICAL_FIELDS))}2"

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    sample_fill = PatternFill("solid", fgColor="FFF4CC")
    for cell in prospects[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for cell in prospects[2]:
        cell.fill = sample_fill
        cell.font = Font(italic=True, color="7C2D12")

    widths = (25, 28, 20, 26, 28, 18, 18, 15, 15, 24, 22, 18, 28, 20, 22, 40)
    for index, width in enumerate(widths, start=1):
        prospects.column_dimensions[get_column_letter(index)].width = width
    prospects.row_dimensions[1].height = 24

    instructions = workbook.create_sheet("Instructions")
    guidance = (
        ("BB Builders Prospect Import Template", ""),
        ("Required", "company_name"),
        ("Recommended", "contact_name; contact_email"),
        ("Optional", "All other columns"),
        ("Rule", "One prospect/company per row"),
        ("Rule", "Do not rename column headers"),
        ("Rule", "Contact email should be a valid email"),
        ("Rule", "Multiple tags can be comma-separated"),
        ("Rule", "Unknown trades may be imported with a warning"),
        ("Rule", "Blank cells are allowed"),
        ("Rule", "Do not use formulas"),
        ("Rule", "Maximum upload size: 2 MB"),
        ("Rule", "Maximum data rows: 1,000"),
        ("Rule", "Supported upload formats: XLSX and CSV"),
        ("Rule", "Delete the example row before importing real data"),
        (
            "Duplicates",
            "Existing companies and contacts may be reused automatically to avoid duplicates.",
        ),
    )
    for row in guidance:
        instructions.append(row)
    instructions["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    instructions["A1"].fill = header_fill
    instructions["B1"].fill = header_fill
    for cell in instructions["A"]:
        if cell.row > 1:
            cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in instructions["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 80

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


HEADER_ALIASES = {
    "company": "company_name",
    "company name": "company_name",
    "business name": "company_name",
    "website": "website",
    "url": "website",
    "phone": "company_phone",
    "company phone": "company_phone",
    "email": "company_email",
    "company email": "company_email",
    "address": "address",
    "city": "city",
    "province": "province",
    "state": "province",
    "province state": "province",
    "postal code": "postal_code",
    "zip": "postal_code",
    "zip code": "postal_code",
    "country": "country",
    "trade": "trade",
    "category": "trade",
    "contact": "contact_name",
    "contact name": "contact_name",
    "title": "contact_title",
    "contact title": "contact_title",
    "contact email": "contact_email",
    "contact phone": "contact_phone",
    "tags": "tags",
    "notes": "notes",
}


def _key(value):
    return " ".join(re.findall(r"[a-z0-9]+", str(value).casefold()))


TRADE_ALIASES = {_key(key): key for key, _ in TRADE_CHOICES}
TRADE_ALIASES.update({_key(label): key for key, label in TRADE_CHOICES})
TRADE_ALIASES.update(
    {
        "mechanical": "hvac-mechanical",
        "hvac": "hvac-mechanical",
        "sprinkler": "fire-protection",
        "sprinklers": "fire-protection",
        "electrical": "electrical-power",
        "data": "low-voltage-data",
        "low voltage": "low-voltage-data",
        "fire protection sprinkler": "fire-protection",
    }
)


def normalize_website(value):
    value = str(value or "").strip()
    if not value:
        return ""
    if "://" not in value:
        value = f"https://{value}"
    URLValidator()(value)
    return value


def normalize_trade(value):
    raw = str(value or "").strip()
    return TRADE_ALIASES.get(_key(raw), ""), raw


def normalize_tags(value):
    items = value if isinstance(value, list) else re.split(r"[,;]", str(value or ""))
    return list(dict.fromkeys(" ".join(item.split()) for item in items if " ".join(item.split())))[
        :20
    ]


def _header(value):
    normalized = _key(value).replace("state province", "province state")
    return HEADER_ALIASES.get(normalized, normalized.replace(" ", "_"))


def _xlsx_rows(content):
    if not content.startswith(b"PK"):
        raise ValidationError("The selected XLSX file is not a valid Excel workbook.")
    try:
        archive = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as error:
        raise ValidationError("The selected XLSX file is not a valid Excel workbook.") from error
    members = archive.infolist()
    if len(members) > 1000 or sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValidationError("The Excel workbook is too large to preview safely.")
    names = {item.filename for item in members}
    if "xl/workbook.xml" not in names:
        raise ValidationError("The selected XLSX file has no readable workbook.")
    shared = []
    if "xl/sharedStrings.xml" in names:
        root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
        shared = ["".join(node.itertext()) for node in root]
    workbook = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    sheet = next((node for node in workbook.iter() if node.tag.endswith("}sheet")), None)
    if sheet is None:
        raise ValidationError("The selected XLSX file has no readable worksheet.")
    relationship_id = next(
        (value for key, value in sheet.attrib.items() if key.endswith("}id")), None
    )
    relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    target = next(
        (
            node.attrib.get("Target", "")
            for node in relationships
            if node.attrib.get("Id") == relationship_id
        ),
        "",
    )
    target = target.lstrip("/")
    sheet_path = target if target.startswith("xl/") else f"xl/{target}"
    if sheet_path not in names:
        raise ValidationError("The first Excel worksheet could not be read.")
    root = ElementTree.fromstring(archive.read(sheet_path))
    rows = []
    for row in (node for node in root.iter() if node.tag.endswith("}row")):
        values = {}
        formula = False
        for cell in (node for node in row if node.tag.endswith("}c")):
            reference = cell.attrib.get("r", "")
            letters = re.match(r"[A-Z]+", reference)
            if not letters:
                continue
            column = 0
            for character in letters.group(0):
                column = column * 26 + ord(character) - 64
            if any(node.tag.endswith("}f") for node in cell):
                formula = True
                value = ""
            elif cell.attrib.get("t") == "inlineStr":
                value = "".join(
                    child.text or "" for child in cell.iter() if child.tag.endswith("}t")
                )
            else:
                value_node = next((child for child in cell if child.tag.endswith("}v")), None)
                value = value_node.text if value_node is not None and value_node.text else ""
                if cell.attrib.get("t") == "s" and value:
                    index = int(value)
                    value = shared[index] if index < len(shared) else ""
            values[column] = str(value)
        width = max(values, default=0)
        rows.append(([values.get(index, "") for index in range(1, width + 1)], formula))
    return rows


def _csv_rows(content):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValidationError("CSV files must use UTF-8 text encoding.") from error
    if "\x00" in text:
        raise ValidationError("The selected CSV file is not valid text data.")
    return [(row, False) for row in csv.reader(io.StringIO(text))]


def _existing_company(organization, values):
    result = ContractorResult(
        display_name=values["company_name"],
        website=values["website"],
        phone=values["company_phone"],
        email=values["company_email"],
        address=values["address"],
        city=values["city"],
        province=values["province"],
        postal_code=values["postal_code"],
        country=values["country"],
        external_place_id="",
    )
    company = dedupe_company(organization, result, "manual_import")
    if company is None and values["company_email"]:
        matches = Company.objects.filter(
            organization=organization, email__iexact=values["company_email"]
        )
        company = matches.first() if matches.count() == 1 else None
    return company


def _existing_contact(company, values):
    email = values.get("contact_email", "")
    if email:
        contact = company.contacts.filter(email__iexact=email).first()
        if contact:
            return contact
    phone = normalize_phone(values.get("contact_phone", ""))
    name = normalized_name(values.get("contact_name", ""))
    matches = [
        contact
        for contact in company.contacts.all()
        if (phone and normalize_phone(contact.phone) == phone)
        or (name and normalized_name(contact.name) == name)
    ]
    return matches[0] if len(matches) == 1 else None


def _clean_values(raw):
    values = {field: str(raw.get(field, "") or "").strip() for field in CANONICAL_FIELDS}
    values["company_name"] = " ".join(values["company_name"].split())
    values["contact_name"] = " ".join(values["contact_name"].split())
    values["country"] = values["country"] or "Canada"
    values["website"] = normalize_website(values["website"])
    values["trade_key"], values["trade_original"] = normalize_trade(values["trade"])
    values["tags_list"] = normalize_tags(raw.get("tags", ""))
    return values


def _analyze_row(organization, row_number, raw, *, formula=False, seen=None):
    errors, warnings = [], []
    try:
        values = _clean_values(raw)
    except ValidationError:
        values = {field: str(raw.get(field, "") or "").strip() for field in CANONICAL_FIELDS}
        values.update(
            {
                "website": str(raw.get("website", "") or "").strip(),
                "country": str(raw.get("country", "") or "").strip() or "Canada",
                "trade_key": "",
                "trade_original": str(raw.get("trade", "") or "").strip(),
                "tags_list": normalize_tags(raw.get("tags", "")),
            }
        )
        errors.append("Website is not valid.")
    if not values["company_name"]:
        errors.append("Company name is required.")
    for field, label in (("company_email", "Company email"), ("contact_email", "Contact email")):
        if values[field]:
            try:
                validate_email(values[field])
                values[field] = values[field].casefold()
            except ValidationError:
                errors.append(f"{label} is not valid.")
    if (
        any(values[field] for field in ("contact_email", "contact_phone", "contact_title"))
        and not values["contact_name"]
    ):
        errors.append("Contact name is required when contact details are supplied.")
    if values["trade_original"] and not values["trade_key"]:
        warnings.append("Trade is not in the controlled taxonomy and will not be assigned.")
    if not values["contact_email"]:
        warnings.append("Contact email is missing.")
    if formula:
        warnings.append("Formula cells were ignored and not executed.")
    fingerprint = "|".join(
        (
            normalize_domain(values["website"]),
            normalize_phone(values["company_phone"]),
            normalized_name(values["company_name"]),
            values["city"].casefold(),
            values["company_email"].casefold(),
            values["contact_email"].casefold(),
            normalized_name(values["contact_name"]),
            normalize_phone(values["contact_phone"]),
        )
    )
    duplicate_in_file = bool(fingerprint.strip("|") and fingerprint in seen)
    if duplicate_in_file:
        warnings.append("Duplicate row in this file; it will be skipped.")
    else:
        seen.add(fingerprint)
    company = _existing_company(organization, values) if not errors else None
    contact = _existing_contact(company, values) if company and values["contact_name"] else None
    if company:
        warnings.append("Existing company will be reused.")
    if contact:
        warnings.append("Existing contact will be reused.")
    status = "invalid" if errors else "warning" if warnings else "valid"
    return {
        "row_number": row_number,
        "status": status,
        "importable": not errors and not duplicate_in_file,
        "errors": errors,
        "warnings": warnings,
        "existing_company_id": company.pk if company else None,
        "existing_contact_id": contact.pk if contact else None,
        "duplicate_in_file": duplicate_in_file,
        "has_valid_contact_email": bool(values["contact_email"] and not errors),
        "values": values,
    }


def preview_import(*, prospect_list, uploaded_file, actor):
    if prospect_list.status != prospect_list.Status.ACTIVE:
        raise ValidationError("Archived prospect lists cannot accept imports.")
    filename = Path(uploaded_file.name).name[:255]
    extension = Path(filename).suffix.casefold()
    if extension not in {".csv", ".xlsx"}:
        raise ValidationError("Choose a CSV or XLSX file.")
    content_type = str(getattr(uploaded_file, "content_type", "") or "").casefold()
    allowed_types = {
        ".csv": {"text/csv", "text/plain", "application/csv", "application/vnd.ms-excel"},
        ".xlsx": {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/zip",
            "application/octet-stream",
        },
    }
    if content_type and content_type not in allowed_types[extension]:
        raise ValidationError("The uploaded content does not match the selected file type.")
    content = uploaded_file.read(MAX_IMPORT_BYTES + 1)
    if not content:
        raise ValidationError("The selected import file is empty.")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValidationError("Import files must be 2 MB or smaller.")
    rows = _csv_rows(content) if extension == ".csv" else _xlsx_rows(content)
    nonblank = [(row, formula) for row, formula in rows if any(str(value).strip() for value in row)]
    if not nonblank:
        raise ValidationError("The import file has no rows.")
    headers = [_header(value) for value in nonblank[0][0]]
    if "company_name" not in headers:
        raise ValidationError("The import file requires a Company Name column.")
    source_rows = nonblank[1:]
    if len(source_rows) > MAX_IMPORT_ROWS:
        raise ValidationError(f"Import files may contain at most {MAX_IMPORT_ROWS} data rows.")
    seen = set()
    preview_rows = []
    for row_number, (row, formula) in enumerate(source_rows, start=2):
        raw = {
            header: row[index] if index < len(row) else ""
            for index, header in enumerate(headers)
            if header in CANONICAL_FIELDS
        }
        if not any(str(value).strip() for value in raw.values()):
            continue
        preview_rows.append(
            _analyze_row(prospect_list.organization, row_number, raw, formula=formula, seen=seen)
        )
    if not preview_rows:
        raise ValidationError("The import file has no non-blank data rows.")
    summary = {
        "total_rows": len(preview_rows),
        "valid_rows": sum(item["status"] == "valid" for item in preview_rows),
        "warning_rows": sum(item["status"] == "warning" for item in preview_rows),
        "invalid_rows": sum(item["status"] == "invalid" for item in preview_rows),
        "likely_duplicates": sum(item["duplicate_in_file"] for item in preview_rows),
        "existing_companies": sum(bool(item["existing_company_id"]) for item in preview_rows),
        "existing_contacts": sum(bool(item["existing_contact_id"]) for item in preview_rows),
        "valid_contact_email": sum(item["has_valid_contact_email"] for item in preview_rows),
        "missing_contact_email": sum(not item["values"]["contact_email"] for item in preview_rows),
    }
    record = ProspectImport.objects.create(
        organization=prospect_list.organization,
        prospect_list=prospect_list,
        filename=filename,
        file_type=extension.removeprefix("."),
        file_digest=hashlib.sha256(content).hexdigest(),
        total_rows=summary["total_rows"],
        valid_rows=summary["valid_rows"],
        warning_rows=summary["warning_rows"],
        error_rows=summary["invalid_rows"],
        preview_rows=preview_rows,
        created_by=actor,
    )
    record_event(
        organization=prospect_list.organization,
        project=None,
        actor=actor,
        action_code="prospect_import.started",
        target=record,
        metadata={"prospect_list_id": prospect_list.pk, "total_rows": record.total_rows},
    )
    return record, summary


def _fill_company(company, values, actor):
    changed = []
    field_map = {
        "website": "website",
        "phone": "company_phone",
        "email": "company_email",
        "address": "address",
        "city": "city",
        "province": "province",
        "postal_code": "postal_code",
        "country": "country",
    }
    for field, source in field_map.items():
        value = values[source]
        if value and not getattr(company, field):
            setattr(company, field, value)
            changed.append(field)
    if changed:
        company.updated_by = actor
        company.save(update_fields=(*changed, "updated_by", "updated_at"))
    return company


def _company_for_values(organization, values, actor, selected_company=None):
    company = selected_company or _existing_company(organization, values)
    if company:
        return _fill_company(company, values, actor), False
    company = Company.objects.create(
        organization=organization,
        display_name=values["company_name"],
        website=values["website"],
        phone=values["company_phone"],
        email=values["company_email"],
        address=values["address"],
        city=values["city"],
        province=values["province"],
        postal_code=values["postal_code"],
        country=values["country"],
        source_type=Company.Source.INTERNAL,
        created_by=actor,
        updated_by=actor,
    )
    return company, True


def _contact_for_values(company, values, actor, *, primary=False):
    if not values["contact_name"]:
        return None, False
    contact = _existing_contact(company, values)
    if contact:
        return contact, False
    return create_contact(
        company=company,
        project=None,
        actor=actor,
        values={
            "name": values["contact_name"],
            "title": values["contact_title"],
            "email": values["contact_email"],
            "phone": values["contact_phone"],
            "is_primary": primary,
            "is_active": True,
        },
    )


def _apply_trade(company, values, *, import_record=None, row_number=None):
    if not values["trade_key"]:
        return None
    capability, _ = TradeCapability.objects.get_or_create(
        company=company,
        trade_key=values["trade_key"],
        defaults={
            "service_cities": [values["city"]] if values["city"] else [],
            "province": values["province"],
            "source_type": company.source_type,
            "source_metadata": {
                "source": "prospect_import" if import_record else "prospect_manual",
                "prospect_import_id": import_record.pk if import_record else None,
                "row_number": row_number,
            },
        },
    )
    return capability


def _apply_tags(entry, values):
    for name in values["tags_list"]:
        tag, _ = ProspectTag.objects.get_or_create(
            organization=entry.prospect_list.organization, name=name
        )
        entry.tags.add(tag)


@transaction.atomic
def add_manual_prospect(*, prospect_list, actor, values):
    if prospect_list.status != prospect_list.Status.ACTIVE:
        raise ValidationError("Archived prospect lists cannot accept new prospects.")
    name = " ".join(str(values.get("name", "")).split())
    email = str(values.get("email", "")).strip().casefold()
    phone = str(values.get("phone", "")).strip()
    company_name = " ".join(str(values.get("company", "")).split())
    validate_email(email)

    email_contacts = list(
        Contact.objects.filter(
            company__organization=prospect_list.organization,
            email__iexact=email,
        ).select_related("company")[:2]
    )
    if company_name:
        named_companies = [
            company
            for company in Company.objects.filter(
                organization=prospect_list.organization, is_active=True
            )
            if normalized_name(company.display_name) == normalized_name(company_name)
            or normalized_name(company.legal_name) == normalized_name(company_name)
        ]
        email_contact = next(
            (
                contact
                for contact in email_contacts
                if contact.company_id in {company.pk for company in named_companies}
            ),
            None,
        )
        if email_contact:
            company = email_contact.company
        elif len(named_companies) == 1:
            company = named_companies[0]
        elif len(named_companies) > 1:
            raise ValidationError(
                {"company": "More than one company matches this name. Use a more specific name."}
            )
        else:
            company = Company.objects.create(
                organization=prospect_list.organization,
                display_name=company_name,
                source_type=Company.Source.INTERNAL,
                created_by=actor,
                updated_by=actor,
            )
    elif len(email_contacts) == 1:
        company = email_contacts[0].company
    elif len(email_contacts) > 1:
        raise ValidationError(
            {"company": "This email matches multiple companies. Enter the company name."}
        )
    else:
        company = Company.objects.create(
            organization=prospect_list.organization,
            display_name=name,
            source_type=Company.Source.INTERNAL,
            created_by=actor,
            updated_by=actor,
        )

    contact = company.contacts.filter(email__iexact=email).first()
    if contact is None:
        contact, _ = create_contact(
            company=company,
            project=None,
            actor=actor,
            values={
                "name": name,
                "title": "",
                "email": email,
                "phone": phone,
                "is_primary": not company.contacts.exists(),
                "is_active": True,
            },
        )
    cleaned = {
        "trade_key": str(values.get("trade", "")),
        "city": company.city,
        "province": company.province,
        "tags_list": normalize_tags(values.get("tags", [])),
        "notes": str(values.get("notes", "")).strip(),
    }
    _apply_trade(company, cleaned)
    status = ProspectListEntry.Status.CONTACT_READY
    entry, created = ProspectListEntry.objects.get_or_create(
        prospect_list=prospect_list,
        company=company,
        defaults={
            "primary_contact": contact,
            "status": status,
            "source_type": "manual",
            "notes": cleaned["notes"],
            "added_by": actor,
        },
    )
    if not created and (
        not entry.is_active or entry.primary_contact_id != contact.pk or entry.status != status
    ):
        entry.is_active = True
        entry.removed_by = None
        entry.removed_at = None
        entry.primary_contact = contact
        entry.status = status
        entry.save(
            update_fields=(
                "is_active",
                "removed_by",
                "removed_at",
                "primary_contact",
                "status",
                "updated_at",
            )
        )
    if cleaned["notes"] and not entry.notes:
        entry.notes = cleaned["notes"]
        entry.save(update_fields=("notes", "updated_at"))
    _apply_tags(entry, cleaned)
    record_event(
        organization=prospect_list.organization,
        project=None,
        actor=actor,
        action_code="prospect.manual_added",
        target=entry,
        metadata={
            "prospect_list_id": prospect_list.pk,
            "company_id": company.pk,
            "entry_created": created,
        },
    )
    return entry, created


@transaction.atomic
def confirm_import(*, import_record, actor):
    import_record = ProspectImport.objects.select_for_update().get(pk=import_record.pk)
    if import_record.status == ProspectImport.Status.COMPLETED:
        return import_record
    imported = skipped = 0
    for row in import_record.preview_rows:
        if not row["importable"]:
            skipped += 1
            continue
        values = row["values"]
        company, _ = _company_for_values(import_record.organization, values, actor)
        contact, _ = _contact_for_values(company, values, actor, primary=False)
        _apply_trade(
            company,
            values,
            import_record=import_record,
            row_number=row["row_number"],
        )
        status = (
            ProspectListEntry.Status.CONTACT_READY
            if contact and contact.is_active and contact.email
            else ProspectListEntry.Status.NEW
        )
        entry, created = ProspectListEntry.objects.get_or_create(
            prospect_list=import_record.prospect_list,
            company=company,
            defaults={
                "primary_contact": contact,
                "status": status,
                "source_type": "spreadsheet_import",
                "source_reference": f"import:{import_record.pk}:row:{row['row_number']}",
                "source_metadata": {
                    "prospect_import_id": import_record.pk,
                    "row_number": row["row_number"],
                    "original_trade": values["trade_original"],
                },
                "notes": values["notes"],
                "added_by": actor,
            },
        )
        if not created:
            if not entry.is_active:
                entry.is_active = True
                entry.removed_by = None
                entry.removed_at = None
                if contact:
                    entry.primary_contact = contact
                    entry.status = status
                entry.save(
                    update_fields=(
                        "is_active",
                        "removed_by",
                        "removed_at",
                        "primary_contact",
                        "status",
                        "updated_at",
                    )
                )
                imported += 1
            else:
                skipped += 1
            if contact and not entry.primary_contact_id:
                entry.primary_contact = contact
                entry.status = status
                entry.save(update_fields=("primary_contact", "status", "updated_at"))
        else:
            imported += 1
        if values["notes"] and not entry.notes:
            entry.notes = values["notes"]
            entry.save(update_fields=("notes", "updated_at"))
        _apply_tags(entry, values)
    import_record.status = ProspectImport.Status.COMPLETED
    import_record.imported_rows = imported
    import_record.skipped_rows = skipped
    import_record.completed_at = timezone.now()
    import_record.save(update_fields=("status", "imported_rows", "skipped_rows", "completed_at"))
    record_event(
        organization=import_record.organization,
        project=None,
        actor=actor,
        action_code="prospect_import.completed",
        target=import_record,
        metadata={
            "prospect_list_id": import_record.prospect_list_id,
            "imported_rows": imported,
            "skipped_rows": skipped,
            "error_rows": import_record.error_rows,
        },
    )
    return import_record
