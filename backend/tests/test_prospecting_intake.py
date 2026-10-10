import io
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.contractors.models import Company, Contact, ScopeContractorCandidate, TradeCapability
from apps.organizations.models import Membership, Organization
from apps.outreach.models import InvitationCampaign, InvitationRecipient
from apps.prospecting.models import ProspectImport, ProspectList, ProspectListEntry

pytestmark = pytest.mark.django_db


def client_for(user):
    client = APIClient()
    client.force_authenticate(user=user)
    return client


def make_list(organization, user):
    return ProspectList.objects.create(
        organization=organization, name="Intake list", created_by=user
    )


def manual_url(organization, prospect_list):
    return (
        f"/api/v1/organizations/{organization.slug}/prospecting/lists/"
        f"{prospect_list.pk}/prospects/manual/"
    )


def imports_url(organization, prospect_list):
    return (
        f"/api/v1/organizations/{organization.slug}/prospecting/lists/{prospect_list.pk}/imports/"
    )


def template_url(organization):
    return f"/api/v1/organizations/{organization.slug}/prospecting/import-template/"


def test_import_template_is_safe_complete_and_available_to_viewers(user, organization, membership):
    before = (
        Company.objects.count(),
        Contact.objects.count(),
        ProspectListEntry.objects.count(),
        ProspectImport.objects.count(),
    )
    membership.role = Membership.Role.VIEWER
    membership.save(update_fields=("role",))

    response = client_for(user).get(template_url(organization))

    assert response.status_code == 200
    assert response["Content-Type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        response["Content-Disposition"]
        == 'attachment; filename="BB-Builders-Prospect-Import-Template.xlsx"'
    )
    workbook = load_workbook(io.BytesIO(b"".join(response.streaming_content)), read_only=False)
    assert workbook.sheetnames == ["Prospects", "Instructions"]
    prospects = workbook["Prospects"]
    assert tuple(cell.value for cell in prospects[1]) == (
        "company_name",
        "website",
        "company_phone",
        "company_email",
        "address",
        "city",
        "province",
        "postal_code",
        "country",
        "trade",
        "contact_name",
        "contact_title",
        "contact_email",
        "contact_phone",
        "tags",
        "notes",
    )
    assert prospects["A2"].value == "Example Mechanical Ltd"
    assert prospects["P2"].value == "Example row - delete before importing"
    assert prospects.freeze_panes == "A2"
    assert "Delete the example row" in " ".join(
        str(cell.value or "") for row in workbook["Instructions"] for cell in row
    )
    assert before == (
        Company.objects.count(),
        Contact.objects.count(),
        ProspectListEntry.objects.count(),
        ProspectImport.objects.count(),
    )


def test_import_template_requires_membership(user, organization, membership):
    other = Organization.objects.create(name="Other template org", slug="other-template-org")
    assert client_for(user).get(template_url(other)).status_code == 403
    assert APIClient().get(template_url(organization)).status_code in {401, 403}


def csv_upload(text):
    return SimpleUploadedFile("prospects.csv", text.encode(), content_type="text/csv")


def xlsx_upload(rows):
    sheet_rows = []
    for row_number, row in enumerate(rows, start=1):
        cells = []
        for column, value in enumerate(row, start=1):
            letters = ""
            number = column
            while number:
                number, remainder = divmod(number - 1, 26)
                letters = chr(65 + remainder) + letters
            cells.append(f'<c r="{letters}{row_number}" t="inlineStr"><is><t>{value}</t></is></c>')
        sheet_rows.append(f'<row r="{row_number}">{"".join(cells)}</row>')
    content = io.BytesIO()
    with zipfile.ZipFile(content, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Prospects" sheetId="1" r:id="rId1"/></sheets></workbook>',
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
            'Target="worksheets/sheet1.xml"/></Relationships>',
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f"<sheetData>{''.join(sheet_rows)}</sheetData></worksheet>",
        )
    return SimpleUploadedFile(
        "prospects.xlsx",
        content.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def test_manual_new_company_contact_trade_and_existing_reuse(user, organization, membership):
    prospect_list = make_list(organization, user)
    payload = {
        "name": "Taylor North",
        "email": "taylor@north.example.com",
        "phone": "807-555-0100",
        "company": "North Mechanical",
        "trade": "hvac-mechanical",
        "tags": ["Northern Ontario"],
    }
    response = client_for(user).post(
        manual_url(organization, prospect_list), payload, format="json"
    )
    assert response.status_code == 201
    company = Company.objects.get(display_name="North Mechanical")
    contact = Contact.objects.get(company=company)
    entry = ProspectListEntry.objects.get(company=company)
    assert contact.email == "taylor@north.example.com"
    assert contact.phone == "807-555-0100"
    assert entry.primary_contact == contact
    assert entry.status == ProspectListEntry.Status.CONTACT_READY
    assert TradeCapability.objects.filter(company=company, trade_key="hvac-mechanical").exists()

    reused = client_for(user).post(manual_url(organization, prospect_list), payload, format="json")
    assert reused.status_code == 200
    assert Company.objects.filter(display_name="North Mechanical").count() == 1
    assert Contact.objects.filter(company=company, email="taylor@north.example.com").count() == 1
    assert ProspectListEntry.objects.filter(company=company).count() == 1


def test_manual_blank_company_creates_minimal_company_and_viewer_is_denied(
    user, organization, membership
):
    prospect_list = make_list(organization, user)
    response = client_for(user).post(
        manual_url(organization, prospect_list),
        {"name": "Jordan Lead", "email": "jordan@example.com"},
        format="json",
    )
    assert response.status_code == 201
    company = Company.objects.get(organization=organization, display_name="Jordan Lead")
    contact = Contact.objects.get(company=company, email="jordan@example.com")
    entry = ProspectListEntry.objects.get(prospect_list=prospect_list, company=company)
    assert entry.primary_contact == contact
    assert entry.status == ProspectListEntry.Status.CONTACT_READY

    membership.role = Membership.Role.VIEWER
    membership.save(update_fields=("role",))
    response = client_for(user).post(
        manual_url(organization, prospect_list),
        {"name": "Denied", "email": "denied@example.com"},
        format="json",
    )
    assert response.status_code == 403


def test_manual_cross_org_access_and_invalid_email_are_rejected(user, organization, membership):
    prospect_list = make_list(organization, user)
    other = Organization.objects.create(name="Other", slug="other-intake")
    other_list = ProspectList.objects.create(
        organization=other,
        name="Other prospects",
        created_by=user,
    )
    response = client_for(user).post(
        manual_url(other, other_list),
        {"name": "Cross Org", "email": "cross@example.com"},
        format="json",
    )
    assert response.status_code == 403

    response = client_for(user).post(
        manual_url(organization, prospect_list),
        {"name": "Bad Email", "email": "not-an-email"},
        format="json",
    )
    assert response.status_code == 400


def test_csv_preview_validates_aliases_duplicates_email_and_unknown_trade(
    user, organization, membership
):
    prospect_list = make_list(organization, user)
    content = (
        "Business Name,URL,City,State,Category,Contact,Contact Email,Tags\n"
        "Alpha HVAC,alpha.example.com,Toronto,ON,HVAC,Alex,alex@alpha.example.com,priority\n"
        "Alpha HVAC,alpha.example.com,Toronto,ON,HVAC,Alex,alex@alpha.example.com,priority\n"
        "Unknown Trade Ltd,,Ottawa,ON,Quantum Trade,Pat,,\n"
        "Bad Email Ltd,,Ottawa,ON,Plumbing,Sam,not-an-email,\n"
    )
    response = client_for(user).post(
        imports_url(organization, prospect_list), {"file": csv_upload(content)}, format="multipart"
    )
    assert response.status_code == 201
    assert response.data["summary"] == {
        "total_rows": 4,
        "valid_rows": 1,
        "warning_rows": 2,
        "invalid_rows": 1,
        "likely_duplicates": 1,
        "existing_companies": 0,
        "existing_contacts": 0,
        "valid_contact_email": 2,
        "missing_contact_email": 1,
    }
    assert ProspectListEntry.objects.count() == 0
    unknown = response.data["rows"][2]
    assert unknown["values"]["trade_key"] == ""
    assert "controlled taxonomy" in unknown["warnings"][0]


def test_xlsx_preview_and_confirm_reuses_company_contact_without_overwrite(
    user, organization, membership
):
    prospect_list = make_list(organization, user)
    company = Company.objects.create(
        organization=organization,
        display_name="Existing Plumbing",
        website="https://existing.example.com",
        email="office@existing.example.com",
        city="Toronto",
        province="ON",
        created_by=user,
        updated_by=user,
    )
    contact = Contact.objects.create(
        company=company, name="Robin Existing", email="robin@existing.example.com"
    )
    upload = xlsx_upload(
        [
            [
                "Company Name",
                "Website",
                "City",
                "Province",
                "Trade",
                "Contact Name",
                "Contact Email",
            ],
            [
                "Existing Plumbing",
                "",
                "Toronto",
                "ON",
                "Plumbing",
                "Robin Existing",
                "robin@existing.example.com",
            ],
        ]
    )
    preview = client_for(user).post(
        imports_url(organization, prospect_list), {"file": upload}, format="multipart"
    )
    assert preview.status_code == 201
    assert preview.data["summary"]["existing_companies"] == 1
    assert preview.data["summary"]["existing_contacts"] == 1
    import_id = preview.data["import"]["id"]
    confirm = client_for(user).post(
        f"{imports_url(organization, prospect_list)}{import_id}/confirm/", {}, format="json"
    )
    assert confirm.status_code == 200
    company.refresh_from_db()
    assert company.website == "https://existing.example.com"
    assert company.email == "office@existing.example.com"
    assert Company.objects.filter(display_name="Existing Plumbing").count() == 1
    assert Contact.objects.filter(company=company).count() == 1
    entry = ProspectListEntry.objects.get(company=company)
    assert entry.primary_contact == contact
    assert entry.status == ProspectListEntry.Status.CONTACT_READY
    assert TradeCapability.objects.filter(company=company, trade_key="plumbing").exists()
    assert ProspectImport.objects.get(pk=import_id).status == ProspectImport.Status.COMPLETED


def test_import_requires_company_header_and_creates_no_procurement_records(
    user, organization, membership
):
    prospect_list = make_list(organization, user)
    missing = client_for(user).post(
        imports_url(organization, prospect_list),
        {"file": csv_upload("City,Contact Email\nToronto,test@example.com\n")},
        format="multipart",
    )
    assert missing.status_code == 400
    preview = client_for(user).post(
        imports_url(organization, prospect_list),
        {
            "file": csv_upload(
                "Company Name,Contact Name,Contact Email\nSafe Import,Casey,casey@example.com\n"
            )
        },
        format="multipart",
    )
    import_id = preview.data["import"]["id"]
    client_for(user).post(
        f"{imports_url(organization, prospect_list)}{import_id}/confirm/", {}, format="json"
    )
    assert InvitationCampaign.objects.count() == 0
    assert InvitationRecipient.objects.count() == 0
    assert ScopeContractorCandidate.objects.count() == 0
